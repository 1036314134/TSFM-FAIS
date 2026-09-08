from __future__ import annotations

import json
import shutil
import sys
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree


OVERVIEW = (
    "图 1 将历史监督路径与部署流程分开，并将部署组件与预注册对照对应起来。"
    "在历史训练起点，教师用可观测未来评估完整候选上下文的预测损失，并从二阶反事实构造缺失块成对效应。"
    "部署时，路由器依次提取极大连续缺失块、生成含安全候选的短名单、利用真实上下文与伪缺失块精炼候选风险，"
    "再通过稀疏关系图、原生有效性约束和确定性束搜索完成块级分配。"
    "选中的块输出被组装为有限数值上下文并交给冻结预测器。"
)

CONTROL_TEXT = (
    "图 1 下方的对照阶梯隔离三项预注册机制变化。"
    "序列重建使用掩码上下文重建误差训练序列选择器；序列预测仅将目标替换为完整候选预测损失；"
    "独立块预测仅改变决策粒度；结构化块预测进一步加入关系图、成对回归器和协调搜索。"
    "最终 B-FAIS 核心对应结构化块预测，并采用 ETT 选择的完整候选目标和纯学习精炼证据。"
)

CAPTION = (
    "图 1　B-FAIS 在冻结 TSFM 预测前，将受损多变量上下文转换为受原生有效性约束的块级补全。"
    "上部显示历史训练起点如何利用完整候选预测损失和二阶反事实效应训练一元排序器与成对模型；"
    "中部以六个编号步骤依次展示缺失块提取、候选短名单、伪缺失证据、风险精炼、联合分配、上下文组装与预测；"
    "下部给出预注册机制检验采用的单因素对照阶梯。"
)


def paragraph_with_text(document: Document, expected: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.strip() == expected]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph {expected!r}, found {len(matches)}")
    return matches[0]


def paragraph_index(document: Document, target: Paragraph) -> int:
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph._p is target._p:
            return index
    raise RuntimeError("Paragraph is not attached to the document")


def replace_text(paragraph: Paragraph, text: str) -> None:
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    paragraph.clear()
    run = paragraph.add_run(text)
    if run_properties is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, run_properties)


def insert_after(paragraph: Paragraph, text: str) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    inserted.style = "Normal"
    inserted.add_run(text)
    return inserted


def first_drawing_after(paragraph: Paragraph) -> Paragraph:
    seen = False
    for candidate in paragraph._parent.paragraphs:
        if candidate._p is paragraph._p:
            seen = True
            continue
        if seen and candidate._p.xpath(".//w:drawing"):
            return candidate
    raise RuntimeError("Could not find the Figure 1 drawing paragraph")


def add_svg_extension(docx_path: Path, svg_path: Path, png_relationship_id: str) -> str:
    content_types_name = "[Content_Types].xml"
    relationships_name = "word/_rels/document.xml.rels"
    document_name = "word/document.xml"
    svg_member_name = "word/media/fig1_overview_zh.svg"

    with zipfile.ZipFile(docx_path, "r") as source_zip:
        files = {item.filename: source_zip.read(item.filename) for item in source_zip.infolist()}

    package_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    relationships = etree.fromstring(files[relationships_name])
    relationship_ids = []
    for relationship in relationships.findall(f"{{{package_ns}}}Relationship"):
        relationship_id = relationship.get("Id", "")
        if relationship_id.startswith("rId") and relationship_id[3:].isdigit():
            relationship_ids.append(int(relationship_id[3:]))
    svg_relationship_id = f"rId{max(relationship_ids, default=0) + 1}"
    relationship = etree.SubElement(relationships, f"{{{package_ns}}}Relationship")
    relationship.set("Id", svg_relationship_id)
    relationship.set(
        "Type",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
    )
    relationship.set("Target", "media/fig1_overview_zh.svg")
    files[relationships_name] = etree.tostring(
        relationships, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )

    content_types_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    content_types = etree.fromstring(files[content_types_name])
    has_svg = any(
        item.get("Extension", "").lower() == "svg"
        for item in content_types.findall(f"{{{content_types_ns}}}Default")
    )
    if not has_svg:
        default = etree.SubElement(content_types, f"{{{content_types_ns}}}Default")
        default.set("Extension", "svg")
        default.set("ContentType", "image/svg+xml")
    files[content_types_name] = etree.tostring(
        content_types, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )

    drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    relationship_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    svg_ns = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
    document = etree.fromstring(files[document_name])
    blips = document.xpath(
        ".//a:blip[@r:embed=$relationship_id]",
        namespaces={"a": drawing_ns, "r": relationship_ns},
        relationship_id=png_relationship_id,
    )
    if len(blips) != 1:
        raise RuntimeError(
            f"Expected one PNG fallback blip for {png_relationship_id}, found {len(blips)}"
        )
    blip = blips[0]
    extension_list = etree.SubElement(blip, f"{{{drawing_ns}}}extLst")
    extension = etree.SubElement(extension_list, f"{{{drawing_ns}}}ext")
    extension.set("uri", "{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
    svg_blip = etree.SubElement(extension, f"{{{svg_ns}}}svgBlip")
    svg_blip.set(f"{{{relationship_ns}}}embed", svg_relationship_id)
    files[document_name] = etree.tostring(
        document, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )
    files[svg_member_name] = svg_path.read_bytes()

    with tempfile.NamedTemporaryFile(
        prefix="figure1-", suffix=".docx", dir=docx_path.parent, delete=False
    ) as temporary_file:
        temporary_path = Path(temporary_file.name)
    try:
        with zipfile.ZipFile(temporary_path, "w", zipfile.ZIP_DEFLATED) as target_zip:
            for name, data in files.items():
                target_zip.writestr(name, data)
        shutil.move(str(temporary_path), str(docx_path))
    finally:
        temporary_path.unlink(missing_ok=True)
    return svg_relationship_id


def main() -> None:
    if len(sys.argv) != 5:
        raise SystemExit("Usage: edit_figure1_docx.py SOURCE OUTPUT PNG SVG")
    source_path = Path(sys.argv[1]).resolve()
    output_path = Path(sys.argv[2]).resolve()
    png_path = Path(sys.argv[3]).resolve()
    svg_path = Path(sys.argv[4]).resolve()

    document = Document(source_path)
    section_heading = paragraph_with_text(document, "4 B-FAIS 与受控变体")
    insert_after(section_heading, OVERVIEW)

    control_heading = paragraph_with_text(document, "4.3 单因素比较阶梯")
    control_index = paragraph_index(document, control_heading)
    control_paragraph = document.paragraphs[control_index + 1]
    if not control_paragraph.text.strip().startswith("图 1 给出相邻对照"):
        raise RuntimeError("The paragraph following Section 4.3 is not the expected overview")
    replace_text(control_paragraph, CONTROL_TEXT)

    drawing_paragraph = first_drawing_after(control_paragraph)
    drawing_index = paragraph_index(document, drawing_paragraph)
    caption_paragraph = document.paragraphs[drawing_index + 1]
    if not caption_paragraph.text.strip().startswith("图 1"):
        raise RuntimeError("The paragraph following the Figure 1 drawing is not its caption")
    replace_text(caption_paragraph, CAPTION)
    caption_paragraph.paragraph_format.keep_with_next = False

    old_width = document.inline_shapes[0].width if document.inline_shapes else None
    section = document.sections[0]
    usable_width = section.page_width - section.left_margin - section.right_margin
    target_width = min(old_width or usable_width, usable_width)

    drawing_paragraph.clear()
    drawing_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    drawing_paragraph.paragraph_format.keep_with_next = True
    picture = drawing_paragraph.add_run().add_picture(str(png_path), width=target_width)
    picture._inline.docPr.set("title", "图 1 B-FAIS 方法总览")
    picture._inline.docPr.set("descr", "B-FAIS 离线监督、六步部署流程与单因素对照阶梯")
    blip = picture._inline.xpath(".//a:blip")[0]
    png_relationship_id = blip.get(qn("r:embed"))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    svg_relationship_id = add_svg_extension(output_path, svg_path, png_relationship_id)

    verification_document = Document(output_path)
    verification_text = "\n".join(paragraph.text for paragraph in verification_document.paragraphs)
    for expected in (OVERVIEW, CONTROL_TEXT, CAPTION):
        if expected not in verification_text:
            raise RuntimeError("A synchronized Figure 1 paragraph is missing after save")
    if len(verification_document.inline_shapes) != 1:
        raise RuntimeError(
            f"Expected one inline figure, found {len(verification_document.inline_shapes)}"
        )

    print(
        json.dumps(
            {
                "output": str(output_path),
                "source_bytes": source_path.stat().st_size,
                "output_bytes": output_path.stat().st_size,
                "figure_width_emu": target_width,
                "png_relationship_id": png_relationship_id,
                "svg_relationship_id": svg_relationship_id,
                "paragraphs": len(verification_document.paragraphs),
                "inline_shapes": len(verification_document.inline_shapes),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
